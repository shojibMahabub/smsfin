"""Excel output writer."""

from pathlib import Path
from typing import Iterable

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill
from openpyxl.utils import get_column_letter
from openpyxl.worksheet.table import Table, TableStyleInfo

from src.models.ledger import LedgerRow, LEDGER_HEADERS


def create_workbook() -> Workbook:
    """Create and return a new Workbook."""
    return Workbook()


def write_sheet(ws, headers: list[str], data_rows: list[list[object]], table_name: str | None = None) -> None:
    """Write data to worksheet with formatting."""
    ws.append(headers)
    for row in data_rows:
        ws.append(row)
    style_header(ws)
    auto_width(ws)
    ws.freeze_panes = "A2"
    ws.auto_filter.ref = ws.dimensions
    if table_name and data_rows:
        ref = f"A1:{get_column_letter(len(headers))}{len(data_rows) + 1}"
        table = Table(displayName=table_name, ref=ref)
        table.tableStyleInfo = TableStyleInfo(name="TableStyleMedium2", showRowStripes=True)
        ws.add_table(table)


def style_header(ws) -> None:
    """Style worksheet header row."""
    fill = PatternFill("solid", fgColor="1F4E78")
    for cell in ws[1]:
        cell.font = Font(bold=True, color="FFFFFF")
        cell.fill = fill


def auto_width(ws) -> None:
    """Auto-adjust column widths."""
    for column in ws.columns:
        max_length = 0
        letter = get_column_letter(column[0].column)
        for cell in column:
            value = "" if cell.value is None else str(cell.value)
            max_length = max(max_length, min(len(value), 60))
        ws.column_dimensions[letter].width = max(10, max_length + 2)


def save_workbook(wb: Workbook, output_path: Path) -> None:
    """Save workbook to file."""
    wb.save(output_path)


def ledger_to_values(rows: list[LedgerRow]) -> list[list[object]]:
    """Convert ledger rows to values for Excel."""
    from src.utils.text import fmt_decimal
    values = []
    for row in rows:
        record = {
            "date": row.date,
            "time": row.time,
            "month": row.month,
            "source": row.source,
            "account": row.account,
            "type": row.type,
            "amount": fmt_decimal(row.amount),
            "currency": row.currency,
            "merchant_or_counterparty": row.merchant_or_counterparty,
            "category": row.category,
            "direction": row.direction,
            "balance_after": fmt_decimal(row.balance_after),
            "transaction_id": row.transaction_id,
            "card_or_account_hint": row.card_or_account_hint,
            "confidence": row.confidence,
            "review_flag": row.review_flag,
            "review_reason": row.review_reason,
            "original_sms": row.original_sms,
        }
        values.append([record[header] for header in LEDGER_HEADERS])
    return values
